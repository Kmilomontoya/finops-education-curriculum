# -*- coding: utf-8 -*-
"""
GENERADOR DE PLANTILLA CURRICULAR FINOPS
Lee de INPUT_DIR:
  - Un Excel con las personas y roles del grupo FinOps (columnas: NOMBRE, ROL, [CARGO])
  - Uno o más CSV del assessment "FinOps Review" de Microsoft (uno por dominio y/o por persona)
Genera en OUTPUT_DIR un Excel con el plan de charlas semanales del FinOps Framework 2026,
priorizado según las brechas derivadas de los CSV del assessment.

Configuración: archivo .env junto al script.
Uso:  python generar_curriculo_finops.py            (menú interactivo)
      python generar_curriculo_finops.py --auto     (ejecución directa)
"""
import csv as csvmod
import re
import sys
import unicodedata
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

BASE = Path(__file__).resolve().parent

# ---------------------------------------------------------------- .env
def cargar_env():
    cfg = {"INPUT_DIR": "input", "OUTPUT_DIR": "output", "ARCHIVO_PERSONAS": "",
           "FECHA_INICIO": "", "NOMBRE_SALIDA": "Plantilla_Curricular_FinOps.xlsx"}
    env = BASE / ".env"
    if env.exists():
        for linea in env.read_text(encoding="utf-8-sig").splitlines():
            linea = linea.strip()
            if not linea or linea.startswith("#") or "=" not in linea:
                continue
            k, v = linea.split("=", 1)
            cfg[k.strip()] = v.strip().strip('"').strip("'")
    return cfg

def normalizar(txt):
    txt = unicodedata.normalize("NFD", str(txt))
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").lower().strip()

# ---------------------------------------------------------------- catálogo del framework 2026
U, Q, O, M, F = ("Understand Usage & Cost", "Quantify Business Value",
                 "Optimize Usage & Cost", "Manage the FinOps Practice", "Fundamentos")
CAPURL = "https://www.finops.org/framework/capabilities/"
DOMURL = "https://www.finops.org/framework/domains/"
# (ID, Tipo, Dominio, Tema, Contenido, NombreMS, URL)
PLAN = [
 ("F-01","Fundamento",F,"¿Qué es FinOps? Definición y Principios","Definición, los 6 principios, valor de negocio de la nube","FinOps Framework","https://www.finops.org/framework/principles/"),
 ("F-02","Fundamento",F,"Personas FinOps","Personas core (Practitioner, Ingeniería, Finanzas, Liderazgo, Producto, Adquisiciones) y aliadas (ITAM, ITFM, ITSM, Seguridad, Sostenibilidad)","","https://www.finops.org/framework/personas/"),
 ("F-03","Fundamento",F,"Fases y Modelo de Madurez","Inform – Optimize – Operate; madurez Crawl / Walk / Run","","https://www.finops.org/framework/phases/"),
 ("F-04","Fundamento",F,"Scopes y Categorías de Tecnología","Segmentos de gasto (producto, centro de costo, ambiente) y categorías: Public Cloud, SaaS, IA, licencias, data center","","https://www.finops.org/framework/scopes/"),
 ("F-05","Fundamento",F,"FOCUS: especificación abierta de costo y uso","Esquema FOCUS, columnas clave y su uso en las exportaciones de Azure Cost Management","","https://focus.finops.org/"),
 ("D-01","Dominio",U,"Dominio: Understand Usage & Cost","Resultados del dominio, sus 4 capabilities y lectura grupal del assessment","Self-Evaluation",DOMURL+"understand-usage-cost/"),
 ("C-01","Capability",U,"Data Ingestion","Fuentes de datos de costo/uso, FOCUS, granularidad y frecuencia de refresco","Data Ingestion",CAPURL+"data-ingestion/"),
 ("C-02","Capability",U,"Allocation","Etiquetado, jerarquías de cuentas, costos compartidos, showback","Allocation",CAPURL+"allocation/"),
 ("C-03","Capability",U,"Reporting & Analytics","Reportes estándar, dashboards, análisis de tendencias y cadencia","Reporting and analytics",CAPURL+"reporting-analytics/"),
 ("C-04","Capability",U,"Anomaly Management","Detección, alertas, triage y respuesta a anomalías de gasto","Anomaly Management",CAPURL+"anomaly-management/"),
 ("D-02","Dominio",Q,"Dominio: Quantify Business Value","Resultados del dominio, sus 5 capabilities y lectura grupal del assessment","Self-Evaluation",DOMURL+"quantify-business-value/"),
 ("C-05","Capability",Q,"Planning & Estimating","Estimación de cargas nuevas y migraciones","Planning and estimating",CAPURL+"planning-estimating/"),
 ("C-06","Capability",Q,"Forecasting","Pronóstico de gasto, métodos y precisión del forecast","Forecasting",CAPURL+"forecasting/"),
 ("C-07","Capability",Q,"Budgeting","Presupuestos, umbrales y control de ejecución","Budgeting",CAPURL+"budgeting/"),
 ("C-08","Capability",Q,"KPIs & Benchmarking","Indicadores FinOps y comparativas internas/externas","Benchmarking",CAPURL+"kpis-benchmarking/"),
 ("C-09","Capability",Q,"Unit Economics","Costo unitario y su relación con el valor de negocio","Unit economics",CAPURL+"unit-economics/"),
 ("D-03","Dominio",O,"Dominio: Optimize Usage & Cost","Resultados del dominio, sus 5 capabilities y lectura grupal del assessment","Self-Evaluation",DOMURL+"optimize-usage-cost/"),
 ("C-10","Capability",O,"Architecting & Workload Placement","Decisiones de arquitectura y ubicación de cargas orientadas a valor","Architecting for cloud",CAPURL+"architecting-workload-placement/"),
 ("C-11","Capability",O,"Usage Optimization","Rightsizing, apagado programado, eliminación de recursos ociosos","Workload optimization",CAPURL+"usage-optimization/"),
 ("C-12","Capability",O,"Rate Optimization","Reservas, savings plans y descuentos negociados","Rate optimization",CAPURL+"rate-optimization/"),
 ("C-13","Capability",O,"Licensing & SaaS","Licenciamiento (p. ej. Azure Hybrid Benefit) y gestión de SaaS","Licensing and SaaS",CAPURL+"licensing-saas/"),
 ("C-14","Capability",O,"Sustainability","Huella de carbono y eficiencia de los recursos en nube","Cloud sustainability",CAPURL+"sustainability/"),
 ("D-04","Dominio",M,"Dominio: Manage the FinOps Practice","Resultados del dominio, sus 8 capabilities y lectura grupal del assessment","Self-Evaluation",DOMURL+"manage-finops-practice/"),
 ("C-15","Capability",M,"Executive Strategy Alignment","Alineación con la estrategia institucional y patrocinio directivo (nueva capability 2026)","",CAPURL+"executive-strategy-alignment/"),
 ("C-16","Capability",M,"FinOps Practice Operations","Operación del equipo FinOps: roles, cadencias, rituales","FinOps practice operations",CAPURL+"finops-practice-operations/"),
 ("C-17","Capability",M,"Governance, Policy & Risk","Políticas, controles (Azure Policy) y gestión de riesgo en nube","Cloud policy and governance",CAPURL+"governance-policy-risk/"),
 ("C-18","Capability",M,"FinOps Education & Enablement","Formación y habilitación — incluye este programa de charlas y sus KPIs","FinOps education and enablement",CAPURL+"finops-education-enablement/"),
 ("C-19","Capability",M,"Invoicing & Chargeback","Facturación, conciliación y cobro interno de costos","Invoicing and chargeback",CAPURL+"invoicing-chargeback/"),
 ("C-20","Capability",M,"FinOps Assessment","Evaluación de madurez, brechas y planes de mejora","FinOps assessment",CAPURL+"finops-assessment/"),
 ("C-21","Capability",M,"Automation, Tools & Services","Automatización, herramientas y servicios de soporte a la práctica","FinOps tools and services",CAPURL+"automation-tools-services/"),
 ("C-22","Capability",M,"Intersecting Disciplines","Intersección con ITAM, ITFM, ITSM, seguridad y sostenibilidad","Intersecting disciplines",CAPURL+"intersecting-disciplines/"),
]
DOMINIOS = [U, Q, O, M]
PRIO_PTS = {"high": 3, "medium": 2, "low": 1}
PERSONA_TOKENS = {
    "Engineering": ["engineering", "ingenieria", "eng"],
    "Finance": ["finance", "financial", "finanzas", "financiero", "financiera"],
    "Leadership": ["leadership", "liderazgo", "lider"],
    "Procurement": ["procurement", "adquisiciones", "compras"],
    "Product": ["product", "producto"],
    "FinOps Practitioner": ["finops", "practitioner", "practicante"],
}

def dominio_de(texto):
    t = normalizar(texto)
    for kw, dom in (("understand", U), ("quantify", Q), ("optimize", O), ("manage", M)):
        if kw in t:
            return dom
    return F

# ---------------------------------------------------------------- lectura de input
def leer_personas(ruta):
    wb = load_workbook(ruta, data_only=True)
    ws = wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        return []
    idx = {"nombre": None, "rol": None, "cargo": None}
    hdr_row = 0
    for i, fila in enumerate(filas[:10]):
        norm = [normalizar(c) if c else "" for c in fila]
        for j, h in enumerate(norm):
            if idx["nombre"] is None and h in ("nombre", "persona", "funcionario", "nombre completo"):
                idx["nombre"] = j; hdr_row = i
            if idx["rol"] is None and h in ("rol", "rol finops", "persona finops", "perfil"):
                idx["rol"] = j; hdr_row = i
            if idx["cargo"] is None and h == "cargo":
                idx["cargo"] = j
        if idx["nombre"] is not None:
            break
    if idx["nombre"] is None:
        idx = {"nombre": 0, "rol": 1, "cargo": 2}
        hdr_row = 0
    personas = []
    for fila in filas[hdr_row + 1:]:
        if not fila or not fila[idx["nombre"]]:
            continue
        personas.append({
            "nombre": str(fila[idx["nombre"]]).strip(),
            "rol": str(fila[idx["rol"]]).strip() if idx["rol"] is not None and len(fila) > idx["rol"] and fila[idx["rol"]] else "",
            "cargo": str(fila[idx["cargo"]]).strip() if idx["cargo"] is not None and len(fila) > idx["cargo"] and fila[idx["cargo"]] else "",
        })
    return personas

def detectar_persona(nombre_archivo, personas):
    n = normalizar(Path(nombre_archivo).stem)
    for p in personas:
        for token in normalizar(p["nombre"]).split():
            if len(token) >= 4 and token in n:
                return p["nombre"]
    for etiqueta, tokens in PERSONA_TOKENS.items():
        if any(t in n for t in tokens):
            return etiqueta
    return "Grupo"

def leer_csv_assessment(ruta, personas):
    """Devuelve (recomendaciones, niveles, persona). recomendaciones: lista de dicts."""
    persona = detectar_persona(ruta.name, personas)
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        filas = list(csvmod.reader(f))
    recomendaciones, niveles = [], {}
    modo_rec, cab = False, []
    for fila in filas:
        if not fila or not any(fila):
            continue
        c0 = fila[0].strip()
        if c0.startswith("-----"):
            modo_rec = False
            continue
        # inicio de la tabla de recomendaciones
        if c0 == "Category" and len(fila) > 1 and fila[1].strip() == "Link-Text":
            cab = [x.strip() for x in fila]
            modo_rec = True
            continue
        # inicio de la tabla de preguntas (fin de recomendaciones)
        if c0 == "Category" and len(fila) > 1 and fila[1].strip() == "Question":
            modo_rec = False
            continue
        # niveles por dominio (solo encabezado: 'Section N: <Dominio>, Nivel, x/y')
        if (not modo_rec and len(fila) >= 3 and c0.lower().startswith("section")
                and re.match(r"^\d+\s*/\s*\d+$", fila[2].strip().strip("'").strip('"'))):
            niveles[dominio_de(c0)] = (fila[1].strip(), fila[2].strip().strip("'").strip('"'))
            continue
        if modo_rec and cab:
            d = dict(zip(cab, fila))
            cap = d.get("Link-Text", "").strip()
            if not cap:
                continue
            try:
                peso = float(str(d.get("Weight", "0")).strip() or 0)
            except ValueError:
                peso = 0.0
            prio = str(d.get("Priority", "")).strip()
            recomendaciones.append({
                "persona": persona,
                "dominio": dominio_de(d.get("Category", "")),
                "capability_ms": cap,
                "prioridad": prio,
                "prio_pts": PRIO_PTS.get(prio.lower(), 0),
                "peso": peso,
                "puntaje": PRIO_PTS.get(prio.lower(), 0) * 100 + peso,
                "contexto": str(d.get("Context", "")).strip(),
            })
    return recomendaciones, niveles, persona

# ---------------------------------------------------------------- agregación y agenda
def agregar(recs):
    """clave (nombre_ms_norm, dominio) -> métricas grupales"""
    grupos = {}
    for r in recs:
        clave = (normalizar(r["capability_ms"]), r["dominio"])
        grupos.setdefault(clave, []).append(r)
    agg = {}
    for clave, lista in grupos.items():
        puntajes = [x["puntaje"] for x in lista]
        prios = [x["prio_pts"] for x in lista]
        agg[clave] = {
            "n": len(lista),
            "puntaje": round(sum(puntajes) / len(puntajes)),
            "prio_prom": sum(prios) / len(prios),
            "divergencia": (max(prios) - min(prios)) if len(lista) >= 2 else None,
            "por_persona": {x["persona"]: x["puntaje"] for x in lista},
            "contexto": max(lista, key=lambda x: x["puntaje"])["contexto"],
        }
    return agg

def prioridad_texto(item):
    if item is None or item["n"] == 0:
        return "SIN DATO"
    p = item["prio_prom"]
    return "ALTA" if p >= 2.5 else ("MEDIA" if p >= 1.5 else "BAJA")

def construir_agenda(agg):
    """Devuelve lista de charlas en orden, con semana asignada y métricas."""
    charlas = []
    for pid, tipo, dom, tema, cont, nms, url in PLAN:
        item = agg.get((normalizar(nms), dom if tipo != "Fundamento" else F)) if nms else None
        # F-01 (FinOps Framework) viene en la sección 1 -> dominio "Fundamentos"
        charlas.append({"id": pid, "tipo": tipo, "dominio": dom, "tema": tema, "contenido": cont,
                        "nombre_ms": nms, "url": url, "agg": item})
    # orden de dominios: mayor brecha promedio primero; sin datos al final en orden canónico
    brecha_dom = {}
    for dom in DOMINIOS:
        pts = [c["agg"]["puntaje"] for c in charlas if c["dominio"] == dom and c["tipo"] == "Capability" and c["agg"]]
        brecha_dom[dom] = sum(pts) / len(pts) if pts else -1
    orden_dom = sorted(DOMINIOS, key=lambda d: (-brecha_dom[d], DOMINIOS.index(d)))
    agenda, semana = [], 1
    for c in charlas:
        if c["tipo"] == "Fundamento":
            c["semana"] = semana; semana += 1
            agenda.append(c)
    for dom in orden_dom:
        intro = next(c for c in charlas if c["tipo"] == "Dominio" and c["dominio"] == dom)
        intro["semana"] = semana; semana += 1
        agenda.append(intro)
        caps = [c for c in charlas if c["tipo"] == "Capability" and c["dominio"] == dom]
        caps.sort(key=lambda c: (-(c["agg"]["puntaje"] if c["agg"] else 0), c["id"]))
        for c in caps:
            c["semana"] = semana; semana += 1
            agenda.append(c)
    # orden sugerido global entre capabilities
    caps = sorted([c for c in agenda if c["tipo"] == "Capability"],
                  key=lambda c: -(c["agg"]["puntaje"] if c["agg"] else 0))
    puntajes = [(c["agg"]["puntaje"] if c["agg"] else 0) for c in caps]
    for c in agenda:
        if c["tipo"] == "Capability":
            p = c["agg"]["puntaje"] if c["agg"] else 0
            c["orden"] = sum(1 for x in puntajes if x > p) + 1
        else:
            c["orden"] = None
    return agenda

# ---------------------------------------------------------------- generación del Excel
AZUL = "1F4E79"; AZUL_CLARO = "DCE6F1"; GRIS = "F2F2F2"
VERDE_F = "C6EFCE"; AMAR_F = "FFEB9C"; ROJO_F = "FFC7CE"; NARANJA_F = "FCE4D6"
FI = Font(name="Arial", size=10); FB = Font(name="Arial", size=10, bold=True)
FH = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FT = Font(name="Arial", size=14, bold=True, color=AZUL)
FINP = Font(name="Arial", size=10, color="0000FF")
FCAL = Font(name="Arial", size=10, color="008000")
FNOTE = Font(name="Arial", size=9, italic=True, color="595959")
HF = PatternFill("solid", start_color=AZUL); SF = PatternFill("solid", start_color=AZUL_CLARO)
thin = Side(style="thin", color="BFBFBF"); BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
CTR = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)

def encabezado(ws, fila, titulos, anchos):
    for i, (h, w) in enumerate(zip(titulos, anchos), 1):
        c = ws.cell(row=fila, column=i, value=h)
        c.font = FH; c.fill = HF; c.border = BRD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w

def generar_excel(agenda, niveles, fuentes, personas, fecha_inicio, ruta_salida):
    wb = Workbook()
    R0 = 5; RN = R0 + len(agenda) - 1
    personas_detectadas = sorted({f["persona"] for f in fuentes}) or ["Grupo"]

    # -------- CURRICULO --------
    wc = wb.create_sheet("CURRICULO")
    wc["A1"] = "PLAN DE CHARLAS SEMANALES — FINOPS FRAMEWORK 2026"; wc["A1"].font = FT
    wc["A2"] = ("Generado automáticamente desde los CSV del assessment. Azul = editable (Semana, Expositor, Estado, Convocados, Asistentes, Notas). "
                "El orden ya viene priorizado por brecha grupal; ajustar Semana si se requiere mover una charla (la fecha se recalcula sola).")
    wc["A2"].font = FNOTE; wc["A2"].alignment = WRAP
    titulos = ["Sem.", "Fecha", "ID", "Tipo", "Dominio", "Charla / Tema", "Contenido sugerido",
               "N.° resp.", "Prioridad grupal", "Puntaje brecha", "Divergencia", "Orden sugerido",
               "Expositor", "Estado", "Convocados", "Asistentes", "% asistencia", "Recomendación destacada del assessment", "Recurso oficial", "Notas"]
    anchos = [7, 13, 7, 12, 24, 32, 44, 8, 13, 11, 11, 10, 22, 13, 11, 11, 11, 55, 36, 24]
    encabezado(wc, 4, titulos, anchos)
    agenda_orden = sorted(agenda, key=lambda c: c["semana"])
    for i, c in enumerate(agenda_orden):
        r = R0 + i
        a = c["agg"]
        wc.cell(row=r, column=1, value=c["semana"]).font = FINP
        wc.cell(row=r, column=2, value=f'=IF($A{r}="","",RESUMEN!$B$4+($A{r}-1)*7)').font = FCAL
        wc.cell(row=r, column=2).number_format = "DD/MM/YYYY"
        wc.cell(row=r, column=3, value=c["id"]).font = FI
        wc.cell(row=r, column=4, value=c["tipo"]).font = FI
        wc.cell(row=r, column=5, value=c["dominio"]).font = FI
        wc.cell(row=r, column=6, value=c["tema"]).font = FB
        wc.cell(row=r, column=7, value=c["contenido"]).font = FI
        wc.cell(row=r, column=8, value=a["n"] if a else ("—" if not c["nombre_ms"] else 0)).font = FCAL
        wc.cell(row=r, column=9, value=("—" if not c["nombre_ms"] else prioridad_texto(a))).font = FCAL
        wc.cell(row=r, column=10, value=(a["puntaje"] if a else ("—" if not c["nombre_ms"] else 0))).font = FCAL
        div = a["divergencia"] if a else None
        wc.cell(row=r, column=11, value=(div if div is not None else "—")).font = FCAL
        wc.cell(row=r, column=12, value=(c["orden"] if c["orden"] else "—")).font = FCAL
        wc.cell(row=r, column=14, value="Programada").font = FINP
        wc.cell(row=r, column=15, value=len(personas) if personas else None).font = FINP
        wc.cell(row=r, column=17, value=f'=IF(OR($O{r}="",$O{r}=0),"",$P{r}/$O{r})').font = FCAL
        wc.cell(row=r, column=17).number_format = "0.0%"
        wc.cell(row=r, column=18, value=(a["contexto"] if a else "")).font = FNOTE
        u = wc.cell(row=r, column=19, value=c["url"]); u.hyperlink = c["url"]
        u.font = Font(name="Arial", size=9, color="0563C1", underline="single")
        for cidx in range(1, 21):
            cel = wc.cell(row=r, column=cidx); cel.border = BRD
            if cidx in (1, 3, 4, 8, 9, 10, 11, 12, 14, 15, 16, 17): cel.alignment = CTR
            if cidx in (7, 18): cel.alignment = WRAP
        for cidx in (13, 16, 20):
            wc.cell(row=r, column=cidx).font = FINP
    dv_est = DataValidation(type="list", formula1='"Programada,Realizada,Aplazada,Cancelada"', allow_blank=True)
    wc.add_data_validation(dv_est); dv_est.add(f"N{R0}:N{RN}")
    if personas:
        dv_exp = DataValidation(type="list", formula1=f"PERSONAS!$A$5:$A${4+len(personas)}", allow_blank=True)
        wc.add_data_validation(dv_exp); dv_exp.add(f"M{R0}:M{RN}")
    for val, fill in (("Realizada", VERDE_F), ("Aplazada", AMAR_F), ("Cancelada", GRIS)):
        wc.conditional_formatting.add(f"N{R0}:N{RN}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", start_color=fill)))
    for val, fill in (("ALTA", ROJO_F), ("MEDIA", NARANJA_F), ("BAJA", AMAR_F)):
        wc.conditional_formatting.add(f"I{R0}:I{RN}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", start_color=fill)))
    wc.conditional_formatting.add(f"K{R0}:K{RN}", CellIsRule(operator="greaterThanOrEqual", formula=["2"], fill=PatternFill("solid", start_color=ROJO_F)))
    wc.freeze_panes = "G5"

    # -------- PERSONAS --------
    wp = wb.create_sheet("PERSONAS")
    wp["A1"] = "GRUPO FINOPS — PERSONAS Y ROLES (desde el Excel de entrada)"; wp["A1"].font = FT
    encabezado(wp, 4, ["Nombre", "Rol FinOps", "Cargo"], [30, 22, 30])
    for i, p in enumerate(personas):
        r = 5 + i
        wp.cell(row=r, column=1, value=p["nombre"]).font = FI
        wp.cell(row=r, column=2, value=p["rol"]).font = FI
        wp.cell(row=r, column=3, value=p["cargo"]).font = FI
        for cidx in range(1, 4):
            wp.cell(row=r, column=cidx).border = BRD

    # -------- FUENTES --------
    wf = wb.create_sheet("FUENTES")
    wf["A1"] = "TRAZABILIDAD — ARCHIVOS DEL ASSESSMENT PROCESADOS"; wf["A1"].font = FT
    encabezado(wf, 4, ["Archivo CSV", "Persona detectada", "Dominios evaluados", "Filas de recomendaciones", "Niveles reportados"], [46, 20, 40, 14, 40])
    for i, fte in enumerate(fuentes):
        r = 5 + i
        wf.cell(row=r, column=1, value=fte["archivo"]).font = FI
        wf.cell(row=r, column=2, value=fte["persona"]).font = FI
        wf.cell(row=r, column=3, value=", ".join(fte["dominios"]) or "—").font = FI
        wf.cell(row=r, column=4, value=fte["n_recs"]).font = FI
        wf.cell(row=r, column=5, value="; ".join(f"{d}: {n} ({p})" for d, (n, p) in fte["niveles"].items()) or "—").font = FNOTE
        for cidx in range(1, 6):
            wf.cell(row=r, column=cidx).border = BRD
            if cidx in (2, 4): wf.cell(row=r, column=cidx).alignment = CTR
        wf.cell(row=r, column=5).alignment = WRAP

    # -------- RESUMEN --------
    wr = wb["Sheet"]; wr.title = "RESUMEN"
    wr["A1"] = "PROGRAMA DE CAPACITACIÓN FINOPS — RESUMEN"; wr["A1"].font = Font(name="Arial", size=16, bold=True, color=AZUL)
    wr["A2"] = f"Generado el {date.today().strftime('%d/%m/%Y')} a partir de {len(fuentes)} archivo(s) de assessment y {len(personas)} persona(s) del grupo."
    wr["A2"].font = FNOTE
    wr["A3"] = "Fecha de corte:"; wr["A3"].font = FB
    wr["B3"] = "=TODAY()"; wr["B3"].number_format = "DD/MM/YYYY"; wr["B3"].font = FB
    wr["A4"] = "Fecha de inicio del ciclo (semana 1):"; wr["A4"].font = FB
    wr["B4"] = fecha_inicio; wr["B4"].number_format = "DD/MM/YYYY"; wr["B4"].font = FINP
    for col, w in zip("ABCDEFG", [46, 16, 14, 14, 14, 14, 30]):
        wr.column_dimensions[col].width = w

    def sec(fila, txt, hasta=7):
        c = wr.cell(row=fila, column=1, value=txt)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF"); c.fill = HF
        for cc in range(2, hasta + 1):
            wr.cell(row=fila, column=cc).fill = HF

    LE = "N"  # columna Estado en CURRICULO
    sec(6, "AVANCE DEL PROGRAMA")
    prog = [
        ("Charlas planificadas", f'=COUNTA(CURRICULO!$C${R0}:$C${RN})', None),
        ("Charlas realizadas", f'=COUNTIF(CURRICULO!${LE}${R0}:${LE}${RN},"Realizada")', None),
        ("% de avance del currículo", '=IF(B7=0,0,B8/B7)', "0.0%"),
        ("Charlas aplazadas", f'=COUNTIF(CURRICULO!${LE}${R0}:${LE}${RN},"Aplazada")', None),
        ("Asistencia promedio (charlas con registro)", f'=IFERROR(AVERAGEIF(CURRICULO!$Q${R0}:$Q${RN},">0"),"—")', "0.0%"),
    ]
    for i, (txt, f_, nf) in enumerate(prog):
        r = 7 + i
        wr.cell(row=r, column=1, value=txt).font = FI
        c = wr.cell(row=r, column=2, value=f_); c.font = FB; c.alignment = CTR
        if nf: c.number_format = nf
        wr.cell(row=r, column=1).border = BRD; c.border = BRD

    sec(13, "NIVEL DE MADUREZ POR DOMINIO (según assessment)")
    for cidx, h in enumerate(["Dominio", "Nivel(es) reportado(s)"], 1):
        c = wr.cell(row=14, column=cidx, value=h); c.font = FB; c.fill = SF; c.border = BRD; c.alignment = CTR
    wr.column_dimensions["B"].width = 40
    for i, dom in enumerate(DOMINIOS):
        r = 15 + i
        vals = niveles.get(dom, [])
        txt = "; ".join(f"{n} ({p})" + (f" — {per}" if per != "Grupo" else "") for per, (n, p) in vals) if vals else "SIN EVALUAR"
        wr.cell(row=r, column=1, value=dom).font = FI
        c = wr.cell(row=r, column=2, value=txt); c.font = FB if vals else FNOTE; c.alignment = WRAP
        if not vals: c.fill = PatternFill("solid", start_color=AMAR_F)
        wr.cell(row=r, column=1).border = BRD; c.border = BRD

    sec(20, "AVANCE POR DOMINIO")
    for cidx, h in enumerate(["Dominio", "Charlas", "Realizadas", "% avance"], 1):
        c = wr.cell(row=21, column=cidx, value=h); c.font = FB; c.fill = SF; c.border = BRD; c.alignment = CTR
    for i, dom in enumerate([F] + DOMINIOS):
        r = 22 + i
        wr.cell(row=r, column=1, value=dom).font = FI
        wr.cell(row=r, column=2, value=f'=COUNTIF(CURRICULO!$E${R0}:$E${RN},$A{r})').font = FCAL
        wr.cell(row=r, column=3, value=f'=COUNTIFS(CURRICULO!$E${R0}:$E${RN},$A{r},CURRICULO!${LE}${R0}:${LE}${RN},"Realizada")').font = FCAL
        wr.cell(row=r, column=4, value=f'=IF(B{r}=0,0,C{r}/B{r})').font = FCAL
        wr.cell(row=r, column=4).number_format = "0.0%"
        for cidx in range(1, 5):
            wr.cell(row=r, column=cidx).border = BRD
            if cidx > 1: wr.cell(row=r, column=cidx).alignment = CTR

    sec(28, "CÓMO USAR ESTE ARCHIVO")
    notas = [
        "1. La agenda de CURRICULO ya viene priorizada: fundamentos primero y luego los dominios/capabilities con mayor brecha según el assessment.",
        "2. Durante el ciclo solo se diligencian las columnas azules: Expositor, Estado (Realizada/Aplazada), Asistentes y Notas. La fecha sale de la Semana y de B4.",
        "3. Divergencia ≥ 2 (en rojo) indica percepciones muy distintas entre quienes respondieron: úsela como punto de discusión en la charla.",
        "4. Al recibir nuevos CSV del assessment, colóquelos en la carpeta input y vuelva a ejecutar el script: se genera un archivo nuevo con la priorización actualizada.",
        "Fuente del marco: FinOps Framework 2026, FinOps Foundation (CC BY 4.0) — finops.org/framework.",
    ]
    for i, n in enumerate(notas):
        wr.cell(row=29 + i, column=1, value=n).font = FNOTE

    wb._sheets = [wb["RESUMEN"], wb["CURRICULO"], wb["PERSONAS"], wb["FUENTES"]]
    wb.save(ruta_salida)

# ---------------------------------------------------------------- flujo principal
def ejecutar(cfg):
    input_dir = (BASE / cfg["INPUT_DIR"]).resolve()
    output_dir = (BASE / cfg["OUTPUT_DIR"]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    if not input_dir.exists():
        print(f"  [ERROR] No existe la carpeta de entrada: {input_dir}")
        return
    # personas
    if cfg["ARCHIVO_PERSONAS"]:
        ruta_personas = input_dir / cfg["ARCHIVO_PERSONAS"]
        xlsx = [ruta_personas] if ruta_personas.exists() else []
    else:
        xlsx = sorted(input_dir.glob("*.xlsx"))
    personas = leer_personas(xlsx[0]) if xlsx else []
    if xlsx:
        print(f"  [OK] Personas: {xlsx[0].name} -> {len(personas)} persona(s)")
    else:
        print("  [AVISO] No se encontró Excel de personas en input/ (se continúa sin roster)")
    # assessments
    csvs = sorted(input_dir.glob("*.csv"))
    if not csvs:
        print("  [ERROR] No se encontraron CSV del assessment en input/")
        return
    recs, fuentes, niveles = [], [], {}
    for ruta in csvs:
        r, niv, per = leer_csv_assessment(ruta, personas)
        recs.extend(r)
        doms = sorted({x["dominio"] for x in r if x["dominio"] != F})
        fuentes.append({"archivo": ruta.name, "persona": per, "dominios": doms, "n_recs": len(r), "niveles": niv})
        for dom, np_ in niv.items():
            niveles.setdefault(dom, []).append((per, np_))
        print(f"  [OK] {ruta.name} -> persona: {per} | dominios: {', '.join(doms) or '—'} | recomendaciones: {len(r)}")
    agg = agregar(recs)
    agenda = construir_agenda(agg)
    # fecha de inicio
    if cfg["FECHA_INICIO"]:
        fi = None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                fi = datetime.strptime(cfg["FECHA_INICIO"], fmt).date(); break
            except ValueError:
                continue
        if fi is None:
            print(f"  [AVISO] FECHA_INICIO '{cfg['FECHA_INICIO']}' no válida; se usa el próximo lunes")
    else:
        fi = None
    if fi is None:
        hoy = date.today()
        fi = hoy + timedelta(days=(7 - hoy.weekday()) % 7 or 7)
    nombre = cfg["NOMBRE_SALIDA"].replace(".xlsx", "") + f"_{date.today().strftime('%Y%m%d')}.xlsx"
    salida = output_dir / nombre
    generar_excel(agenda, niveles, fuentes, personas, fi, salida)
    evaluados = [d for d in DOMINIOS if d in niveles]
    print("  " + "─" * 66)
    print(f"  [OK] Plantilla generada: {salida}")
    print(f"       Charlas: {len(agenda)} | Inicio del ciclo: {fi.strftime('%d/%m/%Y')}")
    print(f"       Dominios con assessment: {len(evaluados)}/4" + (f" (faltan: {', '.join(d for d in DOMINIOS if d not in niveles)})" if len(evaluados) < 4 else ""))

def menu():
    cfg = cargar_env()
    while True:
        print()
        print("  ╔══════════════════════════════════════════════════════════════╗")
        print("  ║        GENERADOR DE PLANTILLA CURRICULAR FINOPS  v1.0          ║")
        print("  ║   Currículo priorizado a partir de un assessment FinOps   ║")
        print("  ╠══════════════════════════════════════════════════════════════╣")
        print(f"  ║  input : {str(cfg['INPUT_DIR']):<52}  ║")
        print(f"  ║  output: {str(cfg['OUTPUT_DIR']):<52}  ║")
        print("  ╠══════════════════════════════════════════════════════════════╣")
        print("  ║  1. Generar plantilla curricular                               ║")
        print("  ║  2. Ver archivos detectados en input                           ║")
        print("  ║  3. Salir                                                      ║")
        print("  ╚══════════════════════════════════════════════════════════════╝")
        op = input("  Opción: ").strip()
        if op == "1":
            ejecutar(cfg)
        elif op == "2":
            d = (BASE / cfg["INPUT_DIR"]).resolve()
            if not d.exists():
                print(f"  [ERROR] No existe: {d}")
                continue
            for f in sorted(d.iterdir()):
                if f.suffix.lower() in (".csv", ".xlsx"):
                    print(f"   - {f.name}")
        elif op == "3":
            break
        else:
            print("  Opción no válida.")

if __name__ == "__main__":
    if "--auto" in sys.argv:
        ejecutar(cargar_env())
    else:
        menu()
